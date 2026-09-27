"""Prepare PAMA passenger car data for Lectures 1 to 3.

The source workbook contains one worksheet per fiscal year and uses presentation
formatting rather than a consistent database layout. This script extracts monthly
production and sales records, normalizes a small set of manufacturer names, and
creates summary tables used by the live demonstration.

The script never modifies the source workbook.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "pama_monthly_production_sales.xlsx"
DATA_DIR = HERE / "data"

MONTH_PATTERN = re.compile(
    r"^(jan|feb|mar|apr|may|jun|june|jul|july|aug|sep|sept|oct|nov|dec)",
    flags=re.IGNORECASE,
)

MANUFACTURER_PATTERNS = [
    ("Honda", re.compile(r"HONDA", re.IGNORECASE)),
    ("Suzuki", re.compile(r"SUZUKI", re.IGNORECASE)),
    ("Toyota", re.compile(r"TOYOTA", re.IGNORECASE)),
    ("Hyundai", re.compile(r"HYUNDAI|SANTRO", re.IGNORECASE)),
    ("BAIC", re.compile(r"BAIC", re.IGNORECASE)),
    ("Honri", re.compile(r"HONRI|DEWAN", re.IGNORECASE)),
    ("Nissan", re.compile(r"NISSAN", re.IGNORECASE)),
    ("Daihatsu", re.compile(r"DAIHATSU|CUORE", re.IGNORECASE)),
    ("Kia", re.compile(r"\bKIA\b", re.IGNORECASE)),
    ("Proton", re.compile(r"PROTON", re.IGNORECASE)),
    ("Chery", re.compile(r"CHERY", re.IGNORECASE)),
    ("United", re.compile(r"UNITED", re.IGNORECASE)),
]


def clean_text(value: object) -> str:
    """Collapse repeated whitespace and return an empty string for blank cells."""

    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def manufacturer_from_label(model_group: str) -> str:
    """Map model labels to a stable manufacturer name."""

    for manufacturer, pattern in MANUFACTURER_PATTERNS:
        if pattern.search(model_group):
            return manufacturer
    return "Other"


def fiscal_months(sheet_name: str) -> list[pd.Timestamp]:
    """Return July to June dates for a worksheet named like 2025-26."""

    start_year = int(sheet_name.split("-")[0])
    return [
        pd.Timestamp(year=start_year if month >= 7 else start_year + 1, month=month, day=1)
        for month in [7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6]
    ]


def locate_passenger_car_block(worksheet) -> tuple[int, int]:
    """Return the first and last row numbers of the passenger car section."""

    start = None
    end = None
    for row in range(1, worksheet.max_row + 1):
        label = clean_text(worksheet.cell(row, 2).value).upper()
        if label == "PASSENGER CARS":
            start = row
            continue
        if start and ("LCVS" in label or label == "TRUCKS"):
            end = row
            break
    if start is None:
        raise ValueError(f"Passenger car section not found in {worksheet.title}")
    return start, end or worksheet.max_row + 1


def locate_month_columns(worksheet, start_row: int) -> list[int]:
    """Find the twelve month columns near the passenger car heading."""

    for row in range(start_row + 1, min(start_row + 12, worksheet.max_row) + 1):
        columns = []
        for column in range(1, worksheet.max_column + 1):
            value = clean_text(worksheet.cell(row, column).value)
            if MONTH_PATTERN.match(value):
                columns.append(column)
        # Some sheets label the cumulative column with a range such as
        # "Jul'08 - June'09". The first twelve month-like cells are the actual
        # months; a later match is the annual summary.
        if len(columns) >= 12:
            return columns[:12]
    raise ValueError(f"Twelve month columns not found in {worksheet.title}")


def source_annual_total(worksheet, row: int, month_columns: list[int]) -> float | None:
    """Read the cached cumulative value immediately after the month columns."""

    for column in range(max(month_columns) + 1, worksheet.max_column + 1):
        value = worksheet.cell(row, column).value
        if isinstance(value, (int, float)):
            return float(value)
    return None


def extract_records() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Extract monthly records and cumulative total validation results."""

    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    records: list[dict] = []
    validation: list[dict] = []

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        start_row, end_row = locate_passenger_car_block(worksheet)
        month_columns = locate_month_columns(worksheet, start_row)
        dates = fiscal_months(sheet_name)
        last_label = ""

        for row in range(start_row + 1, end_row):
            label = clean_text(worksheet.cell(row, 2).value)
            record_type = clean_text(worksheet.cell(row, 3).value).lower().rstrip(".")
            if label:
                last_label = label
            if record_type not in {"prod", "production", "sale", "sales"}:
                continue

            upper_label = last_label.upper()
            if "SUB-TOTAL" in upper_label or "SUB TOTAL" in upper_label or "TOTAL" in upper_label:
                continue

            measure = "Sales" if record_type in {"sale", "sales"} else "Production"
            manufacturer = manufacturer_from_label(last_label)
            monthly_values = [worksheet.cell(row, column).value for column in month_columns]

            if not any(isinstance(value, (int, float)) for value in monthly_values):
                continue

            for date, value in zip(dates, monthly_values):
                records.append(
                    {
                        "fiscal_year": sheet_name,
                        "date": date,
                        "manufacturer": manufacturer,
                        "model_group": last_label,
                        "measure": measure,
                        "units": value if isinstance(value, (int, float)) else pd.NA,
                        "source_sheet": sheet_name,
                        "source_row": row,
                    }
                )

            cached_total = source_annual_total(worksheet, row, month_columns)
            calculated_total = sum(
                float(value) for value in monthly_values if isinstance(value, (int, float))
            )
            validation.append(
                {
                    "fiscal_year": sheet_name,
                    "model_group": last_label,
                    "measure": measure,
                    "source_row": row,
                    "calculated_total": calculated_total,
                    "source_cumulative_total": cached_total,
                    "difference": (
                        calculated_total - cached_total if cached_total is not None else pd.NA
                    ),
                }
            )

    workbook.close()
    return pd.DataFrame(records), pd.DataFrame(validation)


def build_summaries(long_data: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Build annual manufacturer totals and market summaries."""

    sales = long_data.loc[long_data["measure"].eq("Sales")].copy()
    annual_brand = (
        sales.groupby(["fiscal_year", "manufacturer"], as_index=False, sort=False)["units"]
        .sum()
        .rename(columns={"units": "recorded_sales"})
    )

    total_sales = annual_brand.groupby("fiscal_year")["recorded_sales"].sum()
    big_three_sales = (
        annual_brand.loc[annual_brand["manufacturer"].isin(["Suzuki", "Toyota", "Honda"])]
        .groupby("fiscal_year")["recorded_sales"]
        .sum()
    )
    ev_sales = (
        sales.loc[sales["manufacturer"].eq("Honri")]
        .groupby("fiscal_year")["units"]
        .sum()
    )

    market_summary = pd.DataFrame(
        {
            "fiscal_year": list(total_sales.index),
            "total_recorded_sales": total_sales.values,
        }
    )
    market_summary["big_three_recorded_sales"] = market_summary["fiscal_year"].map(
        big_three_sales
    ).fillna(0)
    market_summary["big_three_recorded_share"] = (
        market_summary["big_three_recorded_sales"]
        / market_summary["total_recorded_sales"]
    )
    market_summary["recorded_ev_sales"] = market_summary["fiscal_year"].map(ev_sales).fillna(0)
    market_summary["recorded_ev_share"] = (
        market_summary["recorded_ev_sales"] / market_summary["total_recorded_sales"]
    )
    return annual_brand, market_summary


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Source workbook not found: {SOURCE}")

    DATA_DIR.mkdir(exist_ok=True)
    long_data, validation = extract_records()
    annual_brand, market_summary = build_summaries(long_data)

    long_data.to_csv(DATA_DIR / "passenger_cars_long.csv", index=False, date_format="%Y-%m-%d")
    annual_brand.to_csv(DATA_DIR / "annual_brand_sales.csv", index=False)
    market_summary.to_csv(DATA_DIR / "annual_market_summary.csv", index=False)
    validation.to_csv(DATA_DIR / "source_total_validation.csv", index=False)

    failures = validation.loc[
        validation["source_cumulative_total"].notna()
        & validation["difference"].abs().gt(0.001)
    ]
    sales_rows = long_data.loc[long_data["measure"].eq("Sales")]

    print(f"Source workbook: {SOURCE.name}")
    print(f"Fiscal years: {long_data['fiscal_year'].nunique()}")
    print(f"Monthly records: {len(long_data):,}")
    print(f"Monthly sales records: {len(sales_rows):,}")
    print(f"Validation differences: {len(failures)}")
    print(f"Output directory: {DATA_DIR}")

    if not failures.empty:
        print("\nRows requiring source review:")
        print(failures.to_string(index=False))


if __name__ == "__main__":
    main()
