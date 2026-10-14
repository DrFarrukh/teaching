"""Executed acceptance checks for V4 artifacts (no training or downloads)."""
from io import BytesIO
from pathlib import Path
import hashlib
import json
import posixpath
import re
import zipfile

from lxml import etree
import numpy as np
from PIL import Image
from pptx import Presentation
import torch

CODE=Path(__file__).resolve().parents[1]
ROOT=CODE.parent
OUT=CODE/'v4_outputs'
checks=[]


def check(name,condition,detail=''):
    checks.append({'check':name,'passed':bool(condition),'detail':detail})
    if not condition:raise AssertionError(f'{name}: {detail}')


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    m=json.loads((OUT/'manifest.json').read_text())
    p=ROOT/'AI_Workshop_Session1_v4.pptx'
    prs=Presentation(p)
    sm=json.loads((OUT/'slide_map.json').read_text())
    check('V2 unchanged',sha(ROOT/'AI_Workshop_Session1_v2.pptx')==m['source_v2_sha256'])
    check('Slide count/map',len(prs.slides)==len(sm)==68)
    keys=[r['source_key'] for r in sm]
    check('Requested removals',all(str(n) not in keys for n in (1,6,10,11,37,47)))
    check('Separate pixel-grid slides',keys[keys.index('5'):keys.index('5')+3]==['5','5b','5c'])
    check('Features follow pooling',keys[keys.index('46'):keys.index('46')+3]==['46','44','45'])
    check('Additional photograph examples',all(k in keys for k in ('20b','21b')))
    all_text='\n'.join(sh.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    check('No promotional REAL/Actual wording',not re.search(r'\b(real|actual)\b',all_text,re.I))
    check('No placeholder text',not re.search(r'\b(lorem|ipsum|TODO)\b|\[insert',all_text,re.I))
    bounds=[]
    for i,s in enumerate(prs.slides,1):
        for sh in s.shapes:
            if sh.name.startswith('V4') or sh.shape_type==13:
                if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+20 or sh.top+sh.height>prs.slide_height+20:
                    bounds.append((i,sh.name))
        if i>1:check(f'Page number {i}',any(sh.has_text_frame and sh.text==f'{i} / {len(prs.slides)}' for sh in s.shapes))
    check('Shapes within slide bounds',not bounds,str(bounds))
    # Check every OOXML relationship target resolves, not merely ZIP readability.
    with zipfile.ZipFile(p) as z:
        names=set(z.namelist());bad=[]
        for name in names:
            if not name.endswith('.rels'):continue
            doc=etree.fromstring(z.read(name))
            base=posixpath.dirname(posixpath.dirname(name)) if '/_rels/' in name else ''
            for r in doc:
                if r.get('TargetMode')=='External':continue
                target=r.get('Target')
                resolved=posixpath.normpath(posixpath.join(base,target)) if not target.startswith('/') else target[1:]
                if resolved not in names:bad.append((name,target))
        check('All internal relationships resolve',not bad,str(bad))
        gifs={hashlib.sha256(z.read(n)).hexdigest():z.read(n) for n in names if n.endswith('.gif')}
        check('Four animated GIFs embedded',len(gifs)==4)
        for name in ('convolution.gif','moving_window.gif','relu_pooling.gif','shape_matching.gif'):
            b=(OUT/name).read_bytes();h=hashlib.sha256(b).hexdigest()
            check(f'{name} bytes preserved',h in gifs)
            im=Image.open(BytesIO(b));check(f'{name} animation frames',im.n_frames>=6,str(im.n_frames))
            check(f'{name} loops',im.info.get('loop')==0)
        check('Native editable charts',sum('/charts/chart' in n and n.endswith('.xml') for n in names)>=7)
    c=json.loads((OUT/'evidence/classical.json').read_text())
    x,y,w,h=c['pixel_crop_xywh'];cat=Image.open(CODE/'assets/cat_chelsea.png').convert('L')
    expected=np.asarray(cat)[y:y+h,x:x+w]
    check('Pixel matrix matches source crop',np.array_equal(expected,np.load(OUT/'pixel_values.npy')))
    check('Pixel matrix matches evidence',np.array_equal(expected,c['pixel_values']))
    low=Image.open(CODE/'assets/cat_chelsea.png').convert('RGB').resize((24,16),Image.Resampling.BOX)
    expected_grid=np.asarray(low.resize((720,480),Image.Resampling.NEAREST))
    check('Pixelated cat generated from source',np.array_equal(expected_grid,np.asarray(Image.open(OUT/'cat_pixel_grid.png'))))
    conv=json.loads((OUT/'evidence/convolution.json').read_text())
    inp=np.array(conv['input']);k=np.array(conv['kernel'])
    out=np.array([[np.sum(inp[r:r+3,c:c+3]*k) for c in range(4)] for r in range(4)])
    check('Convolution arithmetic',np.array_equal(out,conv['output']))
    import cv2
    shapes=json.loads((OUT/'evidence/shape_matching.json').read_text())
    edges=np.asarray(Image.open(OUT/'shape_search_edges.png'))
    shape_arrays=np.load(OUT/'shape_matching_responses.npz')
    for name,record in shapes['templates'].items():
        template=np.asarray(Image.open(OUT/f'shape_template_{name}.png'))
        response=cv2.matchTemplate(edges,template,cv2.TM_CCOEFF_NORMED)
        _,score,_,loc=cv2.minMaxLoc(response)
        check(f'{name} complete matching response',np.allclose(response,shape_arrays[name]))
        check(f'{name} best match is algorithm-selected',list(loc)==record['top_left_xy'] and abs(score-record['score'])<1e-6)
    pool=json.loads((OUT/'evidence/pooling.json').read_text());a=np.load(OUT/'feature_activations.npz')
    x,y=pool['source_crop_xy'];ch=pool['source_channel']
    patch=a['relu'][0,ch,y:y+5,x:x+5]
    check('Pooling input is captured model activation',np.allclose(patch,pool['input']))
    calculated=np.array([[patch[r*2:r*2+3,c*2:c*2+3].max() for c in range(2)] for r in range(2)])
    captured=a['maxpool'][0,ch,(y+1)//2:(y+1)//2+2,(x+1)//2:(x+1)//2+2]
    check('Pooling matches model maxpool',np.allclose(calculated,captured) and np.allclose(calculated,pool['output']))
    check('ReLU activation captured correctly',np.array_equal(np.maximum(a['bn1'],0),a['relu']))
    checkpoint=Path(torch.hub.get_dir())/'checkpoints'/m['features']['file']
    state=torch.load(checkpoint,map_location='cpu',weights_only=True)
    vec=torch.from_numpy(a['avgpool'].reshape(1,512))
    with torch.inference_mode():
        logits=vec@state['fc.weight'].T+state['fc.bias']
    error=float(np.max(np.abs(logits.numpy()-a['logits'])))
    check('Pooled features reconstruct classification logits',error<1e-5,f'max error {error:.9g}')
    top=int(torch.softmax(logits,1)[0].argmax())
    check('Top-one feature-flow class is a cat category',281<=top<=285,str(top))
    cm=json.loads((OUT/'evidence/confusion_counts.json').read_text())
    counts=np.array(cm['counts'])
    check('Confusion support is 300 per class',np.array_equal(counts.sum(1),[300]*3))
    check('Confusion accuracy agrees with recorded history',round(counts.trace()/counts.sum()*100,1)==m['lab_history']['ft'][-1][3])
    check('Recorded head example parsed',m['lab_provenance']['head_example']=={'cat':99.7,'dog':.2,'airplane':.1})
    check('Recorded prediction uses its source photograph',sha(OUT/'source/lab_chelsea.png')==m['lab_provenance']['head_example_image_sha256'])
    check('Training not performed',m['training_performed'] is False)
    report={'status':'PASS','checks':checks,'v4_sha256':sha(p),'slide_count':len(prs.slides),
            'presentation_playback':'GIF data verified; playback in PowerPoint/Impress slideshow mode not exercised',
            'visual_review':'Rendered contact sheets inspected by Sol; selected slides inspected at full size.'}
    (OUT/'verification.json').write_text(json.dumps(report,indent=2))
    print(f'PASS: {len(checks)} checks; {len(prs.slides)} slides; 4 embedded animations; editable charts')


if __name__=='__main__':main()
