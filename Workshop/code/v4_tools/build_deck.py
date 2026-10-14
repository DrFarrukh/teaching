"""Build V4 from V2 only. V3 and shared V2/V3 tooling are never modified."""
from pathlib import Path
import hashlib
import json
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from shape_matching import generate as generate_shape_matching

CODE=Path(__file__).resolve().parents[1]
ROOT=CODE.parent
OUT=CODE/'v4_outputs'
ASSETS=CODE/'assets'
SOURCE=OUT/'source'
M=json.loads((OUT/'manifest.json').read_text())
INK='111111';GRAY='525C66';TEAL='087D73';ORANGE='AF5700';PALE='F0F4F6'
prs=Presentation(ROOT/'AI_Workshop_Session1_v2.pptx')
original=list(prs.slides)
mapping=[]


def text(s, content, x, y, w, h, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT, name=None):
    shape=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    shape.name=name or 'V4 text'
    tf=shape.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.alignment=align;p.space_after=Pt(4)
        for r in p.runs:
            r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold
            r.font.color.rgb=RGBColor.from_string(color)
    return shape


def picture(s,path,x,y,w,h,name=None):
    path=Path(path)
    with Image.open(path) as im: iw,ih=im.size
    scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale
    sh=s.shapes.add_picture(str(path),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),
                            width=Inches(dw),height=Inches(dh))
    sh.name=name or path.stem
    # Accessible descriptions without changing appearance.
    sh._element.xpath('.//p:cNvPr')[0].set('descr',name or path.stem.replace('_',' '))
    return sh


def clear(s,section='',title='',subtitle=''):
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    # Removing shapes alone leaves stale media relationships and old GIFs in the package.
    for rid,rel in list(s.part.rels.items()):
        if rel.reltype.endswith('/image') or rel.reltype.endswith('/chart'):
            s.part.drop_rel(rid)
    s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(255,255,255)
    if section:text(s,section,.65,.25,12,.27,12,color=GRAY)
    if title:text(s,title,.65,.69,12.05,.75,32,bold=True,name='V4 title')
    if subtitle:text(s,subtitle,.65,1.43,12.05,.42,18,color=GRAY)


def notes(s,value):
    s.notes_slide.notes_text_frame.text=value


def takeaway(s,value,size=20):
    text(s,value,.65,6.55,12.05,.48,size,bold=True)


def pair(s,left,right,captions=(),top=1.95,height=4.1):
    picture(s,left,.65,top,5.85,height)
    picture(s,right,6.83,top,5.85,height)
    for x,label in zip((.65,6.83),captions):text(s,label,x,top+height+.10,5.85,.36,17,color=GRAY)


def card(s,x,y,w,h,heading,body=''):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(PALE);sh.line.fill.background()
    text(s,heading,x+.20,y+.18,w-.4,.62,23,bold=True)
    if body:text(s,body,x+.20,y+.95,w-.4,h-1.10,18,color=GRAY)


def chart(s,categories,series,x,y,w,h,title,percent=False,legend=True):
    data=CategoryChartData();data.categories=[str(i) for i in categories]
    for name,values in series:data.add_series(name,values)
    c=s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    c.has_title=True;c.chart_title.text_frame.text=title
    for p in c.chart_title.text_frame.paragraphs:
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(19);r.font.bold=True
    c.has_legend=legend
    if legend:
        c.legend.position=XL_LEGEND_POSITION.BOTTOM;c.legend.include_in_layout=False
        c.legend.font.name='Arial';c.legend.font.size=Pt(13)
    for axis in (c.category_axis,c.value_axis):
        axis.tick_labels.font.name='Arial';axis.tick_labels.font.size=Pt(14)
        axis.tick_labels.font.color.rgb=RGBColor.from_string(GRAY)
    c.value_axis.has_major_gridlines=True
    c.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string('E2E7EA')
    c.value_axis.tick_labels.number_format='0' if percent else '0.0'
    c.category_axis.has_title=True;c.category_axis.axis_title.text_frame.text='Epoch'
    for p in c.category_axis.axis_title.text_frame.paragraphs:
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(14)
    for ser,col in zip(c.series,(ORANGE,TEAL)):
        ser.format.line.color.rgb=RGBColor.from_string(col);ser.format.line.width=Pt(2.5)
        ser.marker.size=7
    return c


CM=np.array([[291,8,1],[17,231,52],[7,51,242]])


def confusion(s,x,y,size=4.8):
    # Counts transcribed from the preserved source PNG, not invented or retrained.
    labels=['Airplane','Cat','Dog'];cell=size/4
    text(s,'Predicted →',x+cell,y-.43,size-cell,.32,17,bold=True)
    text(s,'True ↓',x,y-.43,cell,.32,13,color=GRAY)
    for i,label in enumerate(labels):
        text(s,label,x+(i+1)*cell,y,cell,.43,14,bold=True,align=PP_ALIGN.CENTER)
        text(s,label,x,y+(i+1)*cell,cell,cell,14,bold=True,align=PP_ALIGN.CENTER)
    for r in range(3):
        for c in range(3):
            xx=x+(c+1)*cell;yy=y+(r+1)*cell
            sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(xx),Inches(yy),Inches(cell),Inches(cell))
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string('146B65' if r==c else 'E5EFF0')
            sh.line.color.rgb=RGBColor(255,255,255)
            text(s,str(CM[r,c]),xx,yy,cell,cell,30 if size>4 else 23,bold=True,
                 color='FFFFFF' if r==c else INK,align=PP_ALIGN.CENTER)


def lab_curves(s,keys=('scratch',),x=.65,y=1.95,w=12.05,h=4.30):
    hist=M['lab_history'];n=max(len(hist[k]) for k in keys)
    names={'scratch':'Scratch CNN','ft':'Frozen ResNet18 + head'}
    loss=[];acc=[]
    for key in keys:
        a=hist[key]
        loss.append((names[key],[r[1] for r in a]+[None]*(n-len(a))))
        acc.append((names[key],[r[3] for r in a]+[None]*(n-len(a))))
    chart(s,range(1,n+1),loss,x,y,w/2-.18,h,'Training loss',legend=len(keys)>1)
    chart(s,range(1,n+1),acc,x+w/2+.18,y,w/2-.18,h,'Test accuracy (%)',True,legend=len(keys)>1)


def prepare_extra_figures():
    generate_shape_matching()
    a=np.load(OUT/'feature_activations.npz')
    v=a['input'][0]
    rgb=np.clip(v*np.array([.229,.224,.225])[:,None,None]+np.array([.485,.456,.406])[:,None,None],0,1)
    Image.fromarray((rgb.transpose(1,2,0)*255).round().astype(np.uint8)).save(OUT/'feature_model_input.png')
    for layer in ('relu','layer4'):
        values=a[layer][0];chs=M['features']['selected_channels'][layer]
        fig,axes=plt.subplots(2,2,figsize=(6.3,5.8))
        for ax,ch in zip(axes.ravel(),chs):
            ax.imshow(values[ch],cmap='magma',interpolation='nearest');ax.axis('off')
            ax.set_title(f'Channel {ch}',fontsize=16)
        fig.tight_layout();fig.savefig(OUT/f'features_{layer}_grid.png',dpi=150,facecolor='white');plt.close(fig)
    supplied=Image.open(ASSETS/'Segmented_weed.png').convert('RGB')
    supplied.crop((47,30,503,334)).save(SOURCE/'weed_annotation_photo.png')
    supplied.crop((592,30,1048,334)).save(SOURCE/'weed_annotation_mask.png')
    overlay=Image.open(ASSETS/'weed_with_mask.png').convert('RGB')
    overlay.crop((47,31,543,383)).save(OUT/'weed_overlay_clean.png')
    # One explicitly human-defined teaching box annotation, not model output.
    from PIL import ImageDraw
    cat=Image.open(ASSETS/'cat_chelsea.png').convert('RGB')
    ImageDraw.Draw(cat).rectangle((110,35,1195,845),outline='#087D73',width=12)
    cat.save(OUT/'cat_box_annotation.png')
    cm_evidence={'source':'evidence/confusion_matrix_ft.png','source_sha256':hashlib.sha256((OUT/'evidence/confusion_matrix_ft.png').read_bytes()).hexdigest(),
        'counts':CM.tolist(),'transcription':'Visually checked against preserved source PNG',
        'accuracy':float(CM.trace()/CM.sum()),'cat_dog_error_share':float((CM[1,2]+CM[2,1])/(CM.sum()-CM.trace()))}
    (OUT/'evidence/confusion_counts.json').write_text(json.dumps(cm_evidence,indent=2))


def update_slide(key,s):
    if key in (2,3):
        clear(s)
        picture(s,OUT/'cat.png',.85,.45,11.63,5.65)
        text(s,'What do you see?' if key==2 else 'Cat · Fur · Eyes · Whiskers…',.85,6.24,11.63,.63,35,bold=True,align=PP_ALIGN.CENTER)
        notes(s,'Ask the audience before revealing the next slide.' if key==2 else 'Human perception supplies semantic meaning; the next slides reveal numerical representation.')
    elif key==4:
        clear(s,'01 / IMAGE AS DATA','A photograph is a grid of pixels')
        pair(s,OUT/'cat.png',OUT/'cat_pixel_grid.png',('Photograph','24 × 16 samples; nearest-neighbour enlargement'))
        takeaway(s,'Zoom in: each cell is one pixel.')
    elif key in (5,'5b','5c'):
        label={5:'A small patch of the image','5b':'Each pixel has a brightness value','5c':'The image is a matrix of numbers'}[key]
        clear(s,'01 / IMAGE AS DATA',label)
        target={5:'pixels_shades.png','5b':'pixels_shades_numbers.png','5c':'pixels_numbers.png'}[key]
        picture(s,OUT/'pixel_source.png',.65,1.83,6.00,4.5)
        picture(s,OUT/target,7.08,1.83,4.75,4.5)
        takeaway(s,'The same 8 × 8 pixel patch, shown three ways.')
        notes(s,'The three grid slides deliberately remain separate for the presenter to merge. All numbers come from the same original 8x8 crop; coordinates are in evidence/classical.json.')
    elif key==7:
        clear(s,'01 / IMAGE AS DATA','0 is black. 255 is white.','8-bit grayscale has 256 possible intensity levels.')
        for x,value,label in zip((1.05,5.15,9.25),(0,128,255),('Black','Mid-grey','White')):
            sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(2.4),Inches(3),Inches(2.3))
            sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(value,value,value)
            sh.line.color.rgb=RGBColor.from_string('9EA7AE')
            text(s,str(value),x,5.02,3,.7,36,bold=True,align=PP_ALIGN.CENTER)
            text(s,label,x,5.78,3,.4,22,align=PP_ALIGN.CENTER)
    elif key==19:
        clear(s,'02 / CLASSICAL VISION','Contours describe boundaries','Enlarged cat-eye crop: threshold dark pixels, then trace the boundary.')
        pair(s,OUT/'eye_binary_mask.png',OUT/'eye_contour.png',('Thresholded region','Contour of the largest dark region'))
        takeaway(s,'A boundary is geometry—not an object identity.')
    elif key=='20b':
        clear(s,'02 / CLASSICAL VISION','A colour rule applied to a photograph')
        pair(s,ASSETS/'raw_weed.png',OUT/'weed_green_filter.png',('Field photograph','Pixels selected by a green HSV range'))
        takeaway(s,'Green pixels can belong to crops or weeds.')
        notes(s,'HSV range is documented in evidence/classical.json. This is colour thresholding, not a weed recognition model. Lighting and shared colours are limitations.')
    elif key=='21b':
        clear(s,'02 / CLASSICAL VISION','Can simple shapes find the eyes?')
        picture(s,OUT/'shape_matching.gif',.65,1.70,12.05,4.98)
        takeaway(s,'Template → sliding search → similarity map → best match',19)
        notes(s,'Square, circle and oval outlines search the entire Canny edge image at one fixed scale. Animation samples candidate positions; each score comes from the computed normalized-correlation response map. Circle overlaps an eye here; square/oval can select non-eye regions. No eye detector, hand-selected best location or probability of correctness is claimed. A fixed geometric template cannot reliably find all eyes under changes in scale, pose or illumination.')
    elif key==23:
        clear(s,'02 / CLASSICAL VISION','The same scene can look different')
        # Preserve brightness/noise source pictures, replace rotated middle picture.
        for i,path in enumerate((OUT/'cat.png',OUT/'cat_flipped.png',OUT/'cat.png')):
            if i==0:
                a=np.asarray(Image.open(path)).astype(float);Image.fromarray((a*.42).astype(np.uint8)).save(OUT/'cat_dark.png');path=OUT/'cat_dark.png'
            if i==2:
                rng=np.random.default_rng(0);a=np.asarray(Image.open(path)).astype(float)
                Image.fromarray(np.clip(a+rng.normal(0,22,a.shape),0,255).astype(np.uint8)).save(OUT/'cat_noise.png');path=OUT/'cat_noise.png'
            x=.65+i*4.12;picture(s,path,x,2.10,3.8,3.85)
            text(s,('Poor lighting','Horizontal flip','Sensor noise')[i],x,6.12,3.8,.4,20,bold=True)
        notes(s,'The middle photograph is mirrored horizontally, as requested. All panels retain the same scene.')
    elif key==26:
        clear(s)
        picture(s,OUT/'cat.png',.85,.50,11.63,5.60)
        text(s,'Classification: what is in this image?',.85,6.18,11.63,.55,30,bold=True,align=PP_ALIGN.CENTER)
        text(s,'Cat',.85,6.80,11.63,.36,23,align=PP_ALIGN.CENTER)
    elif key==27:
        clear(s,'03 / CLASSIFICATION','A classifier produces class scores','MobileNetV3-Small · ImageNet-1K')
        pair(s,OUT/'cat.png',OUT/'small_cat_scores.png',('Input image','Top three predicted ImageNet categories'))
        takeaway(s,'A score does not supply an object location.')
    elif key==28:
        clear(s,'03 / CLASSIFICATION','Single-label versus multi-label')
        picture(s,OUT/'cat.png',.65,1.95,5.85,3.85)
        picture(s,SOURCE/'bus.jpg',6.83,1.95,5.85,3.85)
        text(s,'Single label: cat',.65,6.00,5.85,.5,24,bold=True)
        text(s,'Multiple image labels: bus + person',6.83,6.00,5.85,.5,24,bold=True)
        takeaway(s,'Multi-label classification reports presence—not boxes or counts.',18)
        notes(s,'These are image-level annotation examples, not claimed multi-label inference. Bus and person are two categories; three people do not create three distinct class labels.')
    elif key==29:
        clear(s,'03 / DETECTION','Detection: what objects—and where?','Faster R-CNN MobileNetV3-320 · score threshold 0.55')
        picture(s,OUT/'faster_bus_boxes.png',1.10,1.9,11.1,4.65)
        takeaway(s,'Class labels + bounding boxes; multiple instances can share a class.')
    elif key==30:
        b=next(b for b in M['inference']['faster']['bus'] if b['label']=='bus')
        clear(s,'03 / DETECTION','One detection, three components')
        picture(s,OUT/'bus_one_box.png',1.10,1.75,11.1,4.60)
        coords=', '.join(str(round(v)) for v in b['box'])
        text(s,f"Bus   ·   score {b['score']:.3f}   ·   [x₁, y₁, x₂, y₂] = [{coords}]",.85,6.42,11.63,.60,22,bold=True,align=PP_ALIGN.CENTER)
        notes(s,'Only the predicted bus box is shown. Coordinates refer to the original image, with origin at the upper-left. A model score is not calibrated proof of correctness.')
    elif key==31:
        clear(s,'03 / DETECTION')
        picture(s,OUT/'iou_boxes.png',1.10,.65,11.1,5.25)
        text(s,'Teal: box A · Orange: box B · Shading: overlap',.85,5.95,11.63,.30,16,color=GRAY,align=PP_ALIGN.CENTER)
        text(s,f"IoU = overlap area / union area = {M['iou']['value']:.2f}",.85,6.30,11.63,.52,29,bold=True,align=PP_ALIGN.CENTER)
        text(s,'Two illustrative boxes; not a ground-truth evaluation',.85,6.90,11.63,.22,14,color=GRAY,align=PP_ALIGN.CENTER)
        notes(s,'Box A and box B are deterministic teaching rectangles, not human ground truth. Their coordinates and executed IoU calculation are preserved in manifest.json.')
    elif key==32:
        clear(s,'03 / SEGMENTATION')
        pair(s,SOURCE/'bus.jpg',OUT/'deeplab_bus_overlay.png',('Input','Predicted class mask'),top=.8,height=5.05)
        text(s,'Segmentation: which pixels belong to each class?',.65,6.30,12.05,.55,28,bold=True,align=PP_ALIGN.CENTER)
        notes(s,'DeepLabV3 MobileNetV3 inference. Bus is red and person is blue; background is left unpainted. These are predicted labels, not annotations.')
    elif key==34:
        clear(s,'03 / SEGMENTATION','Semantic classes versus separate instances')
        paths=sorted(SOURCE.glob('v2_slide_34_*'))
        pair(s,paths[0],paths[1],('One colour per category','Separate colours for separated components'))
        takeaway(s,'Touching objects need an instance-aware model to separate reliably.',18)
        notes(s,'Right panel is connected-component post-processing of a semantic cat mask, not inference from an instance-segmentation network. It illustrates the desired instance distinction only; touching cats could merge.')
        # Existing image text "instances found" would overclaim the method; mask it with a crop.
        for i,path in enumerate(paths):
            im=Image.open(path);im.crop((0,0,im.width,im.height-70)).save(OUT/f'semantic_instance_{i}.png')
        clear(s,'03 / SEGMENTATION','Semantic classes versus separate instances')
        pair(s,OUT/'semantic_instance_0.png',OUT/'semantic_instance_1.png',('Semantic category: cat','Connected components of this mask'))
        takeaway(s,'Illustration of instance separation—not an instance-model prediction.',18)
    elif key==35:
        clear(s,'03 / CORE TASKS','One photograph. Three outputs.')
        for i,path in enumerate((ASSETS/'cat_dog_mix.jpg',OUT/'faster_mix_boxes.png',OUT/'deeplab_mix_overlay.png')):
            x=.65+i*4.12;picture(s,path,x,2.02,3.80,3.8)
            text(s,('Classification: dog','Detection','Segmentation')[i],x,6.02,3.8,.4,22,bold=True)
        takeaway(s,'Whole-image category · Object boxes · Pixel labels')
    elif key==36:
        clear(s,'03 / LABELS','The task determines the annotation')
        for i,path in enumerate((OUT/'cat.png',OUT/'cat_box_annotation.png',SOURCE/'weed_annotation_mask.png')):
            x=.65+i*4.12;picture(s,path,x,2.12,3.8,3.55)
            text(s,('Class label: cat','Bounding-box annotation','Pixel-mask annotation')[i],x,5.96,3.8,.66,21,bold=True)
        takeaway(s,'Training labels are supplied targets—not the model’s predictions.',18)
        notes(s,'Image-level cat label and teaching rectangle are manually defined. Pixel-mask panel is cropped from the user-supplied figure labelled Ground truth (Segmented_weed.png). Class colour names were not supplied; do not guess them.')
    elif key==38:
        clear(s,'03 / APPLICATIONS','From field photograph to segmentation')
        pair(s,ASSETS/'raw_weed.png',OUT/'weed_overlay_clean.png',('Field photograph','Supplied segmentation overlay'))
        takeaway(s,'Pixel labels can support vegetation mapping.')
        notes(s,'User-supplied weed assets. Colours encode the supplied classes; the class-name legend and model provenance were not supplied. Do not report an mIoU from the embedded figure without its evaluation record.')
    elif key in (42,43,46):
        titles={42:'A kernel moves across a pixel grid',43:'The same window moves across the photograph',46:'ReLU and max pooling transform responses'}
        names={42:'convolution.gif',43:'moving_window.gif',46:'relu_pooling.gif'}
        clear(s,'04 / CNN BASICS',titles[key])
        picture(s,OUT/names[key],.65,1.80,12.05,4.82)
        notes(s,('Animated fixed edge-kernel calculation; not a learned CNN kernel. The first frame includes a completed output for static export.' if key!=46 else 'Merged ReLU and pooling. Values come from a pretrained ResNet18 ReLU channel. Actual architecture uses 3x3 max pooling, stride2, padding1; an aligned interior patch is displayed. Negative inputs become zero through ReLU. GIF first frame contains the complete pooled output.'))
    elif key==44:
        clear(s,'04 / LEARNED FEATURES','Feature maps from a pretrained CNN','ResNet18 · same input, different numerical responses')
        picture(s,OUT/'feature_model_input.png',.65,2.1,3.15,3.25)
        picture(s,OUT/'features_relu_grid.png',4.05,1.96,4.10,4.12)
        picture(s,OUT/'features_layer4_grid.png',8.55,1.96,4.1,4.12)
        for x,w,label in ((.65,3.15,'Model input: 224 × 224'),(4.05,4.10,'Early layer: ReLU'),(8.55,4.10,'Deeper layer: layer4')):
            text(s,label,x,6.15,w,.38,18,bold=True)
        text(s,'Selected channels; each display independently scaled for visibility.',.65,6.72,12.05,.31,15,color=GRAY)
        notes(s,'Four channels with highest spatial variance are shown in each layer. These are measured activations, not explanations of semantic concepts. Per-channel min-max scaling prevents intensity comparisons across channels/layers. All tensors and selection IDs are saved.')
    elif key==45:
        clear(s,'04 / CLASSIFICATION','From feature maps to a cat prediction','ResNet18 · ImageNet-1K')
        picture(s,OUT/'feature_model_input.png',.65,2.10,2.85,2.85)
        picture(s,OUT/'features_layer4_grid.png',3.87,1.90,4.1,3.55)
        picture(s,OUT/'features_prediction.png',8.27,2.0,4.45,3.5)
        for x,w,label in ((.65,2.85,'Input image'),(3.87,4.1,'Feature maps → pooled vector'),(8.27,4.45,'Class probabilities')):
            text(s,label,x,5.70,w,.68,20,bold=True)
        takeaway(s,'512 pooled features → learned linear weights → 1,000 class scores',19)
        notes(s,'512-channel layer4 features are global-average-pooled and passed to the ImageNet linear head, then softmax. Selected maps alone do not explain the prediction. Reconstructing logits from the saved pooled vector and head agrees within 1e-5. Cat probability grouping is separate from top-one breed prediction.')
    elif key==48:
        clear(s,'04 / FEATURE INTERPRETATION','Feature maps are responses—not named parts')
        picture(s,OUT/'features_layer1.png',.65,1.95,12.05,3.25)
        card(s,.9,5.53,5.5,1.12,'Different channels respond differently')
        card(s,6.75,5.53,5.5,1.12,'Maps alone do not explain a decision')
        notes(s,'Do not name a channel “eye” or “fur” without evidence. These are selected layer1 activations, independently scaled, with channel IDs saved in the manifest.')
    elif key==51:
        clear(s,'04 / TRAINING','Learning curves help inspect progress')
        lab_curves(s)
        takeaway(s,'Loss and accuracy measure different aspects of learning.',19)
        notes(s,'Historical scratch run: 6000 training images, 900 test images, 5 epochs. Test accuracy was reported each epoch; teach a separate validation set for model selection. Curves are native editable charts using parsed executed-notebook outputs.')
    elif key==52:
        clear(s,'04 / EVALUATION','Separate training, validation and test')
        card(s,.65,2.05,3.75,3.35,'Training','Fit the model’s parameters.')
        card(s,4.8,2.05,3.75,3.35,'Validation','Choose settings and inspect progress.')
        card(s,8.95,2.05,3.75,3.35,'Test','Evaluate once decisions are fixed.')
        takeaway(s,'Do not use test results to choose the model.')
        notes(s,'The supplied historical notebook reports test accuracy each epoch. Repeated reporting alone is not proof of tuning leakage, but a robust workflow separates validation from final testing. No lab code is changed in this deck task.')
    elif key in (55,56,57):
        names={55:'MobileNetV3-Small',56:'SSDLite320 + MobileNetV3-Large',57:'LR-ASPP + MobileNetV3-Large'}
        clear(s,'05 / MODEL EXAMPLES',names[key])
        if key==55:
            picture(s,OUT/'cat.png',.65,1.85,7.0,4.65)
            picture(s,OUT/'small_cat_scores.png',8.0,1.73,4.7,3.5)
            text(s,'Classification · ImageNet-1K',8.0,5.50,4.7,.60,21,bold=True)
            notes(s,'MobileNetV3-Small IMAGENET1K_V1 weights were downloaded with user authorization and verified by checkpoint hash. The chart displays measured top-three class probabilities.')
        elif key==56:
            pair(s,OUT/'ssdlite_mix_boxes.png',OUT/'ssdlite_dog_boxes.png',('Cat–dog photograph','Dog photograph'))
            takeaway(s,'The compact detector can miss or mislabel objects.',19)
            notes(s,'SSDLite COCO_V1, threshold0.55. On this cat–dog photo the model returns two cat-labelled boxes: disclose the wrong label; do not replace the output with Faster R-CNN detections.')
        else:
            pair(s,OUT/'lraspp_cat_overlay.png',OUT/'lraspp_mix_overlay.png',('Cat photograph','Cat–dog photograph'))
            takeaway(s,'One category per pixel, within the supported label set.',19)
            notes(s,'LRASPP COCO_WITH_VOC_LABELS_V1. Purple=cat, green=dog, blue=sofa; background remains unpainted. Compact model outputs are not DeepLab outputs.')
    elif key==59:
        clear(s,'06 / LAB AND MODEL OUTPUTS')
        text(s,'Train a model.\nInspect its outputs.',.85,1.83,11.6,1.6,43,bold=True)
        text(s,'Recorded lab training, then pretrained models on our photographs.',.85,4.03,11.6,.85,25,color=GRAY)
        notes(s,'This section combines existing notebook results with newly executed pretrained inference. Training is not rerun. Model names distinguish the examples from the compact architectures in section05.')
    elif key==60:
        clear(s,'06 / TRAINING','Train a small CNN from scratch','CIFAR-10 subset: airplane · cat · dog')
        lab_curves(s)
        takeaway(s,'5 epochs · 6,000 training images · 900 test images · 71.8% test accuracy',18)
        notes(s,'Historical executed-notebook results. Architecture: three conv/ReLU/max-pool blocks (3→16→32→64 channels), then classifier. Timings omitted because this slide is not a portable speed benchmark.')
    elif key==61:
        clear(s,'06 / TRANSFER LEARNING','Transfer learning performed better here')
        lab_curves(s,('scratch','ft'))
        takeaway(s,'71.8% → 84.9% test accuracy; different models and epoch budgets.',18)
        notes(s,'Historical run: scratch5epochs versus pretrained ResNet18 with frozen backbone and replacement3-class head3epochs. Only1539 headparameters trainable. Useful demonstration, not a controlled claim that transfer learning always wins.')
    elif key==62:
        clear(s,'06 / ERROR ANALYSIS','Where the classifier still errs','Frozen ResNet18 + trained head · 900 test images')
        confusion(s,1.1,2.4,3.8)
        share=(CM[1,2]+CM[2,1])/(CM.sum()-CM.trace())
        text(s,f'{share:.0%}',7.0,2.7,4.6,1.2,65,bold=True,color=TEAL)
        text(s,'of errors confuse cats and dogs',7.0,4.0,4.6,1.1,25,bold=True)
        takeaway(s,'Rows = reference labels. Columns = predicted labels.',19)
    elif key==63:
        r=M['inference']['large']['mix']
        clear(s,'06 / CLASSIFICATION','One top prediction for the whole image','MobileNetV3-Large · ImageNet-1K')
        pair(s,ASSETS/'cat_dog_mix.jpg',OUT/'large_mix_scores.png',('Cat and dog are both present','Probability grouped by ImageNet categories'))
        takeaway(s,f"Dog {r['dog']*100:.1f}% · Cat {r['cat']*100:.1f}% · Other {r['other']*100:.1f}%",22)
        notes(s,'This is a single-label ImageNet classifier; one top prediction is a decision rule, not a limitation of all classifiers. Grouped cat/dog probabilities include their respective ImageNet classes and the remaining “other” probability mass. Size does not guarantee which object wins.')
    elif key==64:
        clear(s,'06 / DETECTION','Confidence is not correctness','Faster R-CNN MobileNetV3-320 · score threshold 0.55')
        pair(s,OUT/'faster_mix_boxes.png',OUT/'faster_fluffy_boxes.png',('Cat and dog found','Dog incorrectly labelled cat'))
        takeaway(s,'A high-scoring prediction can still be wrong.')
    elif key==65:
        clear(s,'06 / SEGMENTATION','Pixel labels can also be wrong','DeepLabV3 + MobileNetV3 · 20 foreground categories + background')
        pair(s,ASSETS/'dog_closeup.jpg',OUT/'deeplab_dog_overlay.png',('Input photograph','Green = dog; orange = predicted chair'))
        takeaway(s,'Part of the fence is mislabelled as chair—not a correct interpretation.',18)
    elif key==66:
        clear(s,'06 / THREE OUTPUTS','One photograph, three different answers')
        for i,path in enumerate((ASSETS/'cat_dog_mix.jpg',OUT/'faster_mix_boxes.png',OUT/'deeplab_mix_overlay.png')):
            x=.65+i*4.12;picture(s,path,x,1.98,3.8,3.65)
            text(s,('Classify: dog','Detect: cat + dog','Segment: pixel categories')[i],x,5.88,3.8,.66,21,bold=True)
        takeaway(s,'Image category · Object locations · Pixel regions')
        notes(s,'Classification: MobileNetV3-Large. Detection: Faster R-CNN MobileNetV3-320. Segmentation: DeepLabV3 MobileNetV3. Each model uses its own weights’ transforms and supported labels; compare output types, not a shared architecture.')
    elif key==68:
        clear(s,'07 / LAB OUTPUTS','Prediction, learning curve, error analysis')
        picture(s,SOURCE/'lab_chelsea.png',.65,2.0,3.3,3.10)
        r=M['lab_provenance']['head_example']
        text(s,f"Cat {r['cat']:.1f}%",.65,5.20,3.3,.64,29,bold=True)
        text(s,'Recorded notebook prediction',.65,5.96,3.3,.48,16,color=GRAY)
        hist=M['lab_history']['ft']
        chart(s,range(1,4),[('ResNet18 head',[v[3] for v in hist])],4.22,2.0,3.8,4.18,'Test accuracy (%)',True,False)
        confusion(s,8.60,2.4,3.75)
        takeaway(s,'All three outputs come from the same recorded ResNet18 lab run.',18)
        notes(s,'Prediction is an external example photograph, not a claimed held-out test example. Values cat99.7%,dog0.2%,airplane0.1% are parsed from executed-notebook stdout. Curve and matrix are from that same trained3-class head. No trained checkpoint was supplied, so no claim of rerunning that head is made.')
    elif key=='takeaways':
        clear(s,'07 / TAKEAWAYS','Choose the output your problem needs')
        card(s,.65,2.02,3.75,3.45,'Pixels are numbers','Preprocessing must match the model.')
        card(s,4.8,2.02,3.75,3.45,'Tasks define outputs','Categories, boxes or pixel labels.')
        card(s,8.95,2.02,3.75,3.45,'Inspect errors','Confidence is not correctness.')
        takeaway(s,'In the lab: build a classifier, evaluate it, then inspect its mistakes.')
    elif key==69:
        clear(s,'FURTHER READING','References and image sources')
        refs=[
            'Szeliski (2022). Computer Vision: Algorithms and Applications, 2nd ed. · szeliski.org/Book',
            'Howard et al. (2019). Searching for MobileNetV3. ICCV.',
            'He et al. (2016). Deep Residual Learning for Image Recognition. CVPR.',
            'Liu et al. (2016). SSD: Single Shot MultiBox Detector. ECCV.',
            'Torchvision model documentation · docs.pytorch.org/vision',
            'Photographs: scikit-image; Unsplash; COCO/Flickr; PyTorch Hub; supplied weed images.',
            'File-level source notes: code/assets/CREDITS.md and code/v4_tools/README.md',
        ]
        for i,line in enumerate(refs):text(s,line,.85,1.90+i*.58,11.60,.47,18,color=GRAY)
        notes(s,'Attribution is based on existing supplied credits, not an independent license audit. Weed image model/class provenance and bus original-photographer details remain to be confirmed before public redistribution.')


def main():
    prepare_extra_figures()
    # Complete all structural changes before editing slide contents.
    order=[2,3,4,5,'5b','5c',7,8,9,12,13,14,15,16,17,18,19,20,'20b',21,'21b',22,23,24,25,
           26,27,28,29,30,31,32,33,34,35,36,38,39,40,41,42,43,46,44,45,48,49,50,51,52,53,
           54,55,56,57,58,59,60,61,62,63,64,65,66,67,68,'takeaways',69]
    bykey={i+1:s for i,s in enumerate(original)}
    blank=min(prs.slide_layouts,key=lambda l:len(l.placeholders))
    for key in order:
        if key not in bykey:bykey[key]=prs.slides.add_slide(blank)
    ids=list(prs.slides._sldIdLst)
    id_by_part={prs.part.related_slide(i.rId).part.partname:i for i in ids}
    selected={bykey[k].part.partname for k in order}
    for sid in ids:
        part=prs.part.related_slide(sid.rId)
        if part.part.partname not in selected:prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    for key in order:prs.slides._sldIdLst.append(id_by_part[bykey[key].part.partname])
    for key in order:update_slide(key,bykey[key])
    # Remove old page numbers and promotional words, preserving run formatting elsewhere.
    for i,(key,s) in enumerate(zip(order,prs.slides),1):
        for sh in list(s.shapes):
            if not sh.has_text_frame:continue
            if re.fullmatch(r'\s*\d+\s*/\s*\d+\s*',sh.text):
                sh._element.getparent().remove(sh._element);continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.text=re.sub(r'\b(?:real|actual)\s+', '',r.text,flags=re.I)
        if i>1:text(s,f'{i} / {len(order)}',11.80,7.16,.88,.18,10,color=GRAY,align=PP_ALIGN.RIGHT,name='V4 page number')
        title=next((sh.text for sh in s.shapes if sh.name=='V4 title'),None)
        if title is None:
            title=' | '.join(sh.text.replace('\n',' ') for sh in s.shapes if sh.has_text_frame)[:130]
        mapping.append({'v4_slide':i,'v2_slide':key if isinstance(key,int) else None,'source_key':str(key),'title':title})
    # Remove stale source section ranges (slide IDs changed); explicit visible section labels remain.
    for ext in list(prs._element.xpath('.//p:extLst')):
        if 'sectionLst' in str(ext.xml):ext.getparent().remove(ext)
    dest=ROOT/'AI_Workshop_Session1_v4.pptx'
    prs.save(dest)
    (OUT/'slide_map.json').write_text(json.dumps(mapping,indent=2))
    (OUT/'slide_map.md').write_text('# V2 → V4 slide map\n\n| V4 | V2/source | Title |\n|---|---|---|\n'+
        '\n'.join(f"| {r['v4_slide']} | {r['source_key']} | {r['title'].replace('|','/')} |" for r in mapping))
    current=hashlib.sha256((ROOT/'AI_Workshop_Session1_v2.pptx').read_bytes()).hexdigest()
    if current!=M['source_v2_sha256']:raise RuntimeError('V2 changed during V4 build; verify source')
    print(f'Built {dest.name}: {len(prs.slides)} slides; V2 hash unchanged')


if __name__=='__main__':main()
